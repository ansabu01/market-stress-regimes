"""
Enhance the seminar deck (dynamic storytelling pass), non-destructively:
  1. TUM logo, top-right, on every framed slide (skips the two full-bleed GIF slides).
  2. Three dark section dividers (matching the title/conclusion aesthetic).
  3. Reorder so each act is introduced by its divider.
Slide transitions are injected in a separate XML pass (add_transitions.py).

Numbers/content are never touched; this only adds branding + structure.
"""
from pathlib import Path
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
SRC = HERE / "final_presentation_draft.pptx"
OUT = HERE / "final_presentation_draft.enhanced.pptx"
LOGO = HERE / "figures" / "tum_logo.png"

# --- palette (verbatim from the existing deck) ---
NAVY   = RGBColor(0x22, 0x26, 0x2E)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
ICE    = RGBColor(0xCA, 0xDC, 0xFC)
CALM   = RGBColor(0x4C, 0x78, 0xA8)
STRESS = RGBColor(0xB2, 0x22, 0x22)
MUTED  = RGBColor(0x9A, 0xA4, 0xB2)   # muted label on dark
GHOST  = RGBColor(0x3A, 0x41, 0x52)   # faint number on dark

EMU_W, EMU_H = None, None

# TUM logo native aspect 1000x521 = 1.919
LOGO_W = Inches(0.62)
LOGO_H = Inches(0.62 / 1.919)   # ~0.323"
LOGO_LEFT = Inches(10.0) - LOGO_W - Inches(0.42)   # right margin 0.42"
LOGO_TOP = Inches(0.26)


def add_logo(slide):
    slide.shapes.add_picture(str(LOGO), LOGO_LEFT, LOGO_TOP, LOGO_W, LOGO_H)


def is_fullbleed_gif(slide, slide_w, slide_h):
    pics = [s for s in slide.shapes if s.shape_type == 13]
    if len(pics) == 1 and not any(s.has_text_frame and s.text_frame.text.strip() for s in slide.shapes):
        p = pics[0]
        if p.width and p.width >= slide_w * 0.98 and p.height >= slide_h * 0.98:
            return True
    return False


def set_dark_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY


def add_dots(slide, left_in, top_in, size_in=0.09):
    for i, col in enumerate((CALM, STRESS)):
        sq = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                    Inches(left_in + i * (size_in + 0.055)), Inches(top_in),
                                    Inches(size_in), Inches(size_in))
        sq.fill.solid(); sq.fill.fore_color.rgb = col
        sq.line.fill.background()
        sq.shadow.inherit = False


def txt(slide, l, t, w, h, text, *, font="Calibri", size=14, color=WHITE, bold=False,
        italic=False, align=PP_ALIGN.LEFT, spacing=None, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color
    if spacing is not None:
        rPr = r._r.get_or_add_rPr(); rPr.set("spc", str(int(spacing)))
    return tb


def make_divider(prs, number, part_label, title, subtitle, accent):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    set_dark_bg(slide)
    # eyebrow: dots + PART label
    add_dots(slide, 0.55, 1.30, 0.10)
    txt(slide, 0.87, 1.19, 6.0, 0.34, part_label, font="Calibri", size=11,
        color=MUTED, bold=True, spacing=260)
    # hero number
    txt(slide, 0.52, 1.66, 2.2, 1.7, number, font="Cambria", size=96, color=accent, bold=True)
    # title + subtitle to the right of the number
    txt(slide, 2.30, 2.02, 7.0, 1.0, title, font="Cambria", size=30, color=WHITE, bold=True)
    txt(slide, 2.32, 3.16, 6.9, 0.7, subtitle, font="Calibri", size=14.5, color=MUTED)
    add_logo(slide)
    return slide


def main():
    prs = Presentation(str(SRC))
    sw, sh = prs.slide_width, prs.slide_height

    # 1) logo on all existing slides except full-bleed GIF slides
    gif_positions = []
    for idx, slide in enumerate(prs.slides):
        if is_fullbleed_gif(slide, sw, sh):
            gif_positions.append(idx)
            continue
        add_logo(slide)
    print("full-bleed GIF slides (0-idx, logo skipped):", gif_positions)

    # 2) three dark section dividers (created at end; repositioned below)
    d1 = make_divider(prs, "01", "PART 01  ·  THE QUESTION",
                      "Does diversification survive when it is needed most?",
                      "Why a full-sample correlation is the wrong lens for crisis risk.", ICE)
    d2 = make_divider(prs, "02", "PART 02  ·  DATING THE REGIME",
                      "A data-driven calm-versus-stress split of the market",
                      "A two-state Markov-switching model on monthly SPY returns.", ICE)
    d3 = make_divider(prs, "03", "PART 03  ·  WHAT BREAKS, AND WHAT HOLDS",
                      "Correlations, concentration, and the volatility correction",
                      "From pairwise co-movement to portfolio-level fragility.", ICE)

    # 3) reorder sldIdLst: divider before Motivation(1), Regime(3), Results(7)
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)                       # [orig0..orig15, d1, d2, d3]
    orig = ids[:16]
    D1, D2, D3 = ids[16], ids[17], ids[18]
    new_order = (
        [orig[0], D1, orig[1], orig[2],        # Title, D1, Motivation, Data
         D2, orig[3], orig[4], orig[5], orig[6],  # D2, regime static, chain, gif5, gif6
         D3] + orig[7:]                         # D3, correlations..backups
    )
    for el in ids:
        sldIdLst.remove(el)
    for el in new_order:
        sldIdLst.append(el)

    prs.save(str(OUT))
    print(f"saved -> {OUT.name}  ({len(list(prs.slides))} slides)")


if __name__ == "__main__":
    main()
