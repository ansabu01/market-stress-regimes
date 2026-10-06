"""Inject smooth Fade slide transitions (schema-correct position) into every slide.
Dividers + title + conclusion get a slower fade to read as act-breaks; content slides
a medium fade. Purely presentational; touches no content."""
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

HERE = Path(__file__).resolve().parent
DECK = HERE / "final_presentation_draft.enhanced.pptx"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

# 0-indexed slides that are dark "beat" slides (title, 3 dividers, conclusion) -> slow fade
SLOW = {0, 1, 4, 9, 16}


def make_transition(speed):
    t = etree.SubElement(etree.Element(qn("p:sld")), qn("p:transition"))
    t.set("spd", speed)
    etree.SubElement(t, qn("p:fade"))
    return t


def main():
    prs = Presentation(str(DECK))
    for i, slide in enumerate(prs.slides):
        sld = slide._element
        # remove any existing transition
        for old in sld.findall(qn("p:transition")):
            sld.remove(old)
        speed = "slow" if i in SLOW else "med"
        trans = make_transition(speed)
        clr = sld.find(qn("p:clrMapOvr"))
        csld = sld.find(qn("p:cSld"))
        if clr is not None:
            clr.addnext(trans)
        elif csld is not None:
            csld.addnext(trans)
        else:
            sld.append(trans)
    prs.save(str(DECK))
    print(f"transitions added to {len(list(prs.slides))} slides (slow on {sorted(SLOW)})")


if __name__ == "__main__":
    main()
